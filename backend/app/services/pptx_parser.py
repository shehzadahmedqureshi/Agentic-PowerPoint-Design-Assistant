from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas import ElementFeature, ElementType, PresentationAnalysis, SlideStructure


def _ratio(value: int, total: int) -> float:
    return round(max(0.0, min(1.0, value / total)), 4)


def _shape_type(shape) -> ElementType:
    if getattr(shape, "has_chart", False):
        return ElementType.CHART
    if getattr(shape, "has_table", False):
        return ElementType.TABLE
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return ElementType.IMAGE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return ElementType.GROUP
    if getattr(shape, "has_text_frame", False):
        return ElementType.TEXT_BOX
    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return ElementType.SHAPE
    return ElementType.OTHER


def _font_metadata(shape) -> tuple[float | None, bool | None, str | None, int | None]:
    if not getattr(shape, "has_text_frame", False):
        return None, None, None, None

    text = shape.text or ""
    sizes: list[float] = []
    bold_values: list[bool] = []
    alignment = None
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.alignment is not None and alignment is None:
            alignment = str(paragraph.alignment).split(" ")[0].lower()
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(round(run.font.size.pt, 1))
            if run.font.bold is not None:
                bold_values.append(run.font.bold)
    return (
        round(sum(sizes) / len(sizes), 1) if sizes else None,
        any(bold_values) if bold_values else None,
        alignment,
        len(text),
    )


def _fill_color(shape) -> str | None:
    try:
        if shape.fill.type is None or shape.fill.fore_color.type is None:
            return None
        rgb = shape.fill.fore_color.rgb
        return str(rgb) if rgb else None
    except (AttributeError, TypeError, ValueError):
        return None


def _table_dimensions(shape) -> tuple[int | None, int | None]:
    if not getattr(shape, "has_table", False):
        return None, None
    return len(shape.table.rows), len(shape.table.columns)


def _describe(elements: list[ElementFeature]) -> str:
    counts = {kind: 0 for kind in ElementType}
    for element in elements:
        counts[element.element_type] += 1

    parts = [
        "Landscape presentation slide",
        f"with {len(elements)} visual elements",
        f"including {counts[ElementType.TEXT_BOX]} text boxes",
        f"{counts[ElementType.IMAGE]} images",
        f"{counts[ElementType.CHART]} charts",
        f"{counts[ElementType.TABLE]} tables",
        f"and {counts[ElementType.SHAPE]} decorative shapes.",
    ]

    wide = [e for e in elements if e.width >= 0.55]
    right_images = [e for e in elements if e.element_type == ElementType.IMAGE and e.x >= 0.5]
    top_text = [e for e in elements if e.element_type == ElementType.TEXT_BOX and e.y <= 0.2]
    if top_text:
        parts.append("Text is positioned near the top, consistent with a title region.")
    if right_images:
        parts.append("A prominent image region appears on the right side.")
    if len(wide) >= 2:
        parts.append("Multiple elements span most of the slide width.")

    for index, element in enumerate(elements[:12], start=1):
        detail = (
            f"Element {index} is a {element.element_type.value} at "
            f"({element.x:.2f}, {element.y:.2f}) sized "
            f"({element.width:.2f}, {element.height:.2f})"
        )
        if element.font_size is not None:
            detail += f" with average font size {element.font_size:.1f}pt"
        if element.character_count is not None:
            detail += f" and {element.character_count} characters"
        if element.table_rows is not None and element.table_columns is not None:
            detail += (
                f" with {element.table_rows} rows and "
                f"{element.table_columns} columns"
            )
        parts.append(detail + ".")
    return " ".join(parts)


def analyze_presentation(path: Path) -> PresentationAnalysis:
    presentation = Presentation(path)
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    slides: list[SlideStructure] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        elements: list[ElementFeature] = []
        for shape in slide.shapes:
            font_size, bold, alignment, character_count = _font_metadata(shape)
            table_rows, table_columns = _table_dimensions(shape)
            elements.append(
                ElementFeature(
                    element_type=_shape_type(shape),
                    x=_ratio(shape.left, slide_width),
                    y=_ratio(shape.top, slide_height),
                    width=_ratio(shape.width, slide_width),
                    height=_ratio(shape.height, slide_height),
                    font_size=font_size,
                    bold=bold,
                    alignment=alignment,
                    fill_color=_fill_color(shape),
                    character_count=character_count,
                    table_rows=table_rows,
                    table_columns=table_columns,
                )
            )
        slides.append(
            SlideStructure(
                slide_number=slide_number,
                width=slide_width,
                height=slide_height,
                elements=elements,
                structural_description=_describe(elements),
            )
        )

    return PresentationAnalysis(
        filename=path.name,
        slide_count=len(slides),
        slides=slides,
    )
