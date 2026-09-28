import shutil
import uuid
from pathlib import Path

from pptx import Presentation
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models import TemplateSlide
from app.services.embeddings import embed_descriptions
from app.services.pptx_parser import analyze_presentation
from app.services.previews import render_pptx_previews


def _excluded_slide_numbers(path: Path) -> set[int]:
    presentation = Presentation(path)
    excluded: set[int] = set()
    for slide_number, slide in enumerate(presentation.slides, start=1):
        local_text = " ".join(
            shape.text.lower()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text
        )
        if "credits" in local_text and (
            "free for everyone" in local_text or "thanks to" in local_text
        ):
            excluded.add(slide_number)
    return excluded


async def ingest_reference_deck(session: AsyncSession, path: Path) -> int:
    settings = get_settings()
    analysis = analyze_presentation(path)
    embeddings = embed_descriptions(
        [slide.structural_description for slide in analysis.slides]
    )
    excluded = _excluded_slide_numbers(path)

    await session.execute(
        delete(TemplateSlide).where(TemplateSlide.source_filename == path.name)
    )
    render_work = settings.previews_dir / "render-work"
    shutil.rmtree(render_work, ignore_errors=True)
    rendered = render_pptx_previews(path, render_work)
    ingested = 0
    for slide, embedding, rendered_image in zip(
        analysis.slides, embeddings, rendered, strict=True
    ):
        if slide.slide_number in excluded:
            rendered_image.unlink(missing_ok=True)
            continue
        template_id = uuid.uuid4()
        preview_path = settings.previews_dir / f"{template_id}.png"
        shutil.move(rendered_image, preview_path)
        session.add(
            TemplateSlide(
                id=template_id,
                name=f"{path.stem} — Slide {slide.slide_number}",
                source_filename=path.name,
                slide_number=slide.slide_number,
                preview_path=str(preview_path.resolve()),
                structure=slide.model_dump(mode="json"),
                structural_description=slide.structural_description,
                structural_embedding=embedding,
            )
        )
        ingested += 1
    await session.commit()
    shutil.rmtree(render_work, ignore_errors=True)
    return ingested


async def template_count(session: AsyncSession) -> int:
    result = await session.scalar(select(TemplateSlide.id).limit(1))
    return 1 if result else 0
