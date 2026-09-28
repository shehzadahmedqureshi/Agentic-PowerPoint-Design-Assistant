from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.schemas import ElementType
from app.services.pptx_parser import analyze_presentation


def test_analyze_presentation_does_not_expose_raw_text(tmp_path: Path) -> None:
    source = tmp_path / "sample.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Revenue increased by 35%"
    run.font.size = Pt(28)
    deck.save(source)

    result = analyze_presentation(source)

    assert result.slide_count == 1
    element = result.slides[0].elements[0]
    assert element.element_type == ElementType.TEXT_BOX
    assert element.character_count == 24
    assert element.font_size == 28
    assert "Revenue" not in result.slides[0].structural_description
    assert "24 characters" in result.slides[0].structural_description


def test_analysis_filename_can_be_replaced_with_original_upload_name(tmp_path: Path) -> None:
    stored_file = tmp_path / "internal-id.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(stored_file)

    result = analyze_presentation(stored_file)
    result.filename = "Quarterly Results.pptx"

    assert result.filename == "Quarterly Results.pptx"
